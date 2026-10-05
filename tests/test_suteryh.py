"""Encargados de edificio (CCT 589/10): motor, escalas y flujo web.

Casos sobre la planilla de SUTERH de septiembre 2026. Los marcados "recibo real"
reproducen al centavo recibos de consorcios de Bahía Blanca, sin datos personales.
"""
import json
import re
import sqlite3
from datetime import date
from decimal import Decimal as D
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from backend import db as dbmod
from backend.app import create_app
from backend.calculo_suteryh import (ADICIONALES, CATEGORIAS_SUTERYH, liquidar_sac_suteryh, liquidar_suteryh,
                                     liquidar_zona_fria, nombres_escala)
from backend.escalas import generar_plantilla

PLANILLA = json.loads((Path(__file__).parent.parent / "backend" / "datos" /
                       "escalas_suteryh_2026_jul_sep.json").read_text(encoding="utf-8"))
SUTERYH = "CCT 589/10"
PAGO = {"fecha_pago": "2026-10-03", "lugar_pago": "Bahía Blanca", "ultimo_deposito_periodo": "08/2026",
        "ultimo_deposito_fecha": "2026-09-10", "ultimo_deposito_banco": "Banco Nación"}


def liquidar(periodo, cargo, categoria, anios, **kw):
    tabla = PLANILLA["periodos"][periodo]
    datos = tabla["cargos"][cargo]
    anio, mes = (int(p) for p in periodo.split("-"))
    return liquidar_suteryh(
        periodo=periodo, cargo=datos["nombre"], categoria_edificio=categoria,
        basico=D(str(datos["basico_por_categoria"].get(str(categoria), 1))),
        adicionales={k: D(str(v)) for k, v in tabla["adicionales"].items()},
        vigencia_escala=date(anio, mes, 1), fecha_ingreso=date(anio - anios, mes, 1), **kw)


def importes(liq):
    return {c.codigo: c.importe for c in liq.conceptos}


# --- Motor -----------------------------------------------------------------------
# Los casos con "recibo real" reproducen al centavo recibos de septiembre 2026 de
# consorcios de Bahía Blanca, con los datos personales sacados (quedan fuera del repo).

def test_vigilancia_nocturna_no_afiliado_y_zona_fria_aparte():
    """Recibo real: vigilancia nocturna, 11 años, sin cuota sindical; zona fría en recibo aparte."""
    liq = liquidar("2026-09", "vigilancia_nocturna", 3, 11, afiliado=False)
    assert importes(liq) == {
        "BAS": D("1294224.00"), "ANT": D("266313.30"), "ADR": D("85000.00"), "JUB": D("181009.10"),
        "PAMI": D("49366.12"), "OS": D("49366.12"), "CPF": D("16455.37"), "FMVDD": D("16455.37"),
        "A27B": D("12341.53"), "RED": D("0.31")}
    assert liq.total_remunerativo == D("1645537.30") and liq.neto == D("1320544")
    # Contribuciones del convenio del recibo: CAPAF 4%, FMVDD 1,5%, art. 27 bis 0,75%, SERACARH 0,5%.
    assert {c.codigo: c.importe for c in liq.informativos} == {
        "C_CPF": D("65821.49"), "C_FMVDD": D("24683.06"), "C_A27B": D("12341.53"), "C_SERACARH": D("8227.69")}

    zona = liquidar_zona_fria(liq, porcentaje=D("50"), afiliado=False)
    assert importes(zona) == {
        "ZONA": D("822768.65"), "JUB": D("90504.55"), "PAMI": D("24683.06"), "OS": D("24683.06"),
        "CPF": D("8227.69"), "FMVDD": D("8227.69"), "A27B": D("6170.76"), "RED": D("0.16")}
    assert zona.neto == D("660272")


def test_encargado_permanente_con_residuos_y_tareas():
    """Recibo real: encargado sin vivienda, 3ª categoría, 19 años, 38 UF, cocheras, jardín y viáticos."""
    liq = liquidar("2026-09", "encargado_permanente_sv", 3, 19, unidades_funcionales=38,
                   tareas=("limpieza_cocheras", "jardin", "viaticos"))
    assert importes(liq) == {
        "BAS": D("1331567.00"), "ADR": D("85000.00"), "ANT": D("459995.70"), "RES": D("90022.00"),
        "COCH": D("32328.60"), "JARD": D("32328.60"), "VIAT": D("90035.30"), "JUB": D("233340.49"),
        "PAMI": D("63638.32"), "OS": D("63638.32"), "CPF": D("21212.77"), "FMVDD": D("21212.77"),
        "A27B": D("15909.58"), "SIND": D("42425.54"), "RED": D("0.59")}
    assert liq.total_remunerativo == D("2121277.20")
    # Zona fría de ese consorcio: 50% de básico + antigüedad, en recibo aparte.
    zona = liquidar_zona_fria(liq, porcentaje=D("50"), zona_base="basico_antiguedad")
    assert importes(zona)["ZONA"] == D("895781.35") and zona.neto == D("700949")


def test_horas_extra_al_100_sin_viaticos_ni_suma_fija():
    """Recibo real: 6 hs al 100%, valor hora = (básico + antig. + residuos + cocheras + jardín) / 200."""
    liq = liquidar("2026-09", "encargado_permanente_sv", 4, 9, unidades_funcionales=42, horas_100=6,
                   tareas=("limpieza_cocheras", "jardin", "viaticos"))
    i = importes(liq)
    assert i["HE100"] == D("95553.77")
    assert liq.total_remunerativo == D("1863151.97") and liq.neto == D("1457917")


def test_horas_extra_al_50():
    liq = liquidar("2026-09", "encargado_permanente_sv", 4, 0, horas_50=10, horas_100=4)
    # valor hora = 1.210.515 / 200 = 6.052,575 (la suma fija no entra)
    assert importes(liq)["HE50"] == D("90788.63")
    assert importes(liq)["HE100"] == D("48420.60")


def test_encargado_no_permanente_jornada_reducida():
    """Recibo real: no permanente sin vivienda, 9 años: mitad de suma fija y antigüedad al 1% exacto."""
    liq = liquidar("2026-09", "encargado_no_permanente_sv", 1, 9, unidades_funcionales=17,
                   tareas=("limpieza_cocheras", "viaticos"))
    i = importes(liq)
    assert (i["BAS"], i["ADR"], i["ANT"], i["RES"]) == (D("761130.00"), D("42500.00"), D("108946.35"),
                                                      D("40273.00"))
    assert liq.total_remunerativo == D("1075213.25") and liq.neto == D("841355")
    assert liq.jornada_horas == 4
    zona = liquidar_zona_fria(liq, porcentaje=D("50"), zona_base="basico_antiguedad")
    assert importes(zona)["ZONA"] == D("435038.18") and zona.neto == D("340418")


def test_horas_extra_jornada_reducida_divisor_100():
    liq = liquidar("2026-09", "encargado_no_permanente_cv", 1, 0, horas_100=8)
    assert importes(liq)["HE100"] == D("114349.76")  # 714.686 / 100 × 2 × 8


def test_no_afiliado_no_paga_cuota_sindical():
    liq = liquidar("2026-09", "encargado_permanente_sv", 1, 0, afiliado=False)
    assert "SIND" not in importes(liq) and "VIV" not in importes(liq)
    assert liq.total_remunerativo == D("1537618.00")
    assert liq.total_descuentos == D("303679.56")
    assert liq.neto == D("1233939")


def test_alicuota_sindical_local():
    liq = liquidar("2026-09", "encargado_permanente_sv", 1, 0, alicuota_sindical=D("0.025"))
    assert importes(liq)["SIND"] == D("38440.45")


def test_media_jornada_antiguedad_1pct_y_mitad_de_adicional():
    liq = liquidar("2026-09", "ayudante_media_jornada", 3, 5)
    i = importes(liq)
    assert (i["BAS"], i["ADR"], i["ANT"]) == (D("665783.00"), D("42500.00"), D("60525.75"))
    assert liq.jornada_horas == 4


def test_vivienda_solo_si_se_pide():
    """Los recibos reales no liquidan vivienda ni en cargos "con vivienda"; el art. 15 queda como opción."""
    assert "VIV" not in importes(liquidar("2026-09", "encargado_permanente_cv", 1, 0))
    liq = liquidar("2026-09", "encargado_permanente_cv", 1, 0, dias=15, liquidar_vivienda=True)
    i = importes(liq)
    assert (i["BAS"], i["ADR"], i["VIV"], i["VIVE"]) == (D("633899.50"), D("42500.00"), D("4351.55"),
                                                         D("4351.55"))


def test_titulo_tareas_y_zona_en_el_mismo_recibo():
    liq = liquidar("2026-09", "encargado_permanente_sv", 1, 0, tramos_titulo=2,
                   tareas=("jardin", "limpieza_piletas", "limpieza_cocheras"))
    i = importes(liq)
    assert i["TIT"] == D("145261.80")  # 10% de 1.452.618
    assert (i["JARD"], i["PILE"], i["COCH"]) == (D("32328.60"), D("54384.60"), D("32328.60"))
    zona = liquidar("2026-09", "encargado_permanente_sv", 1, 0, zona_desfavorable=True, horas_50=10)
    # 50% de todo lo remunerativo, horas extra incluidas: (1.452.618 + 85.000 + 108.946,35) / 2
    assert importes(zona)["ZONA"] == D("823282.18")


def test_errores_del_motor():
    with pytest.raises(ValueError, match="tarea desconocida"):
        liquidar("2026-09", "encargado_permanente_sv", 1, 0, tareas=("lavar autos",))
    with pytest.raises(ValueError, match="categoría del edificio"):
        liquidar("2026-09", "encargado_permanente_sv", 5, 0)
    with pytest.raises(ValueError, match="base de zona"):
        liquidar("2026-09", "encargado_permanente_sv", 1, 0, zona_base="todo")
    with pytest.raises(ValueError, match="falta en la escala: Valor vivienda"):
        liquidar_suteryh(periodo="2026-09", cargo="Encargado Permanente con vivienda", categoria_edificio=1,
                         basico=D(1), adicionales={"adicional_remuneratorio_mensual": D(1)},
                         vigencia_escala=date(2026, 9, 1), fecha_ingreso=date(2026, 1, 1), liquidar_vivienda=True)


def test_sac_con_aportes_del_convenio():
    meses = [liquidar(p, "encargado_permanente_cv", 2, 10, unidades_funcionales=20)
             for p in ("2026-07", "2026-08", "2026-09")]
    liq = liquidar_sac_suteryh(periodo="2026-09", categoria="Encargado Permanente con vivienda",
                               fecha_ingreso=date(2016, 9, 1), fecha_egreso=date(2026, 9, 30), historial=meses)
    i = importes(liq)
    assert i["SAC"] == D("397364.25")  # 1.589.457 / 2 × 92/184
    assert "VIVE" not in i
    assert (i["JUB"], i["A27B"], i["SIND"]) == (D("43710.07"), D("2980.23"), D("7947.29"))
    assert liq.neto == D("310938")


@pytest.mark.parametrize("periodo", ["2026-07", "2026-08", "2026-09"])
def test_plus_antiguedad_es_2pct_del_ayudante_sv_4ta(periodo):
    tabla = PLANILLA["periodos"][periodo]
    base = D(str(tabla["cargos"]["ayudante_permanente_sv"]["basico_por_categoria"]["4"]))
    assert abs(base * D("0.02") - D(str(tabla["adicionales"]["plus_antiguedad_2pct"]))) < D("0.1")


def test_cargos_del_codigo_coinciden_con_la_planilla():
    for tabla in PLANILLA["periodos"].values():
        assert [c["nombre"] for c in tabla["cargos"].values()] == CATEGORIAS_SUTERYH
        assert set(tabla["adicionales"]) == set(ADICIONALES)


# --- Base, escalas y pantallas ---------------------------------------------------------
@pytest.fixture
def client(tmp_path):
    return create_app({"DATABASE": str(tmp_path / "s.db"), "TESTING": True}).test_client()


def consorcio(client, **edificio) -> int:
    r = client.post("/empresas/nueva", data={"razon_social": "Consorcio Calle Falsa 123", "cuit": "30712345671",
                                             "domicilio": "Calle Falsa 123", "lugar_pago": "Bahía Blanca"})
    e = int(r.headers["Location"].rstrip("/").split("/")[-1])
    if edificio:
        r = client.post(f"/empresas/{e}/edificio", data=edificio, follow_redirects=True)
        assert "Datos del edificio guardados" in r.text
    return e


def encargado(client, e, **campos):
    datos = {"apellido": "Pérez", "nombre": "Ana", "cuil": "27-33333333-9", "convenio": SUTERYH,
             "categoria": "Encargado Permanente con vivienda", "fecha_ingreso": "2016-09-01",
             "afiliado": "1", "retira_residuos": "1", "tramos_titulo": "0", **campos}
    return client.post(f"/empresas/{e}/empleados", data=datos, follow_redirects=True)


def test_instalacion_trae_escalas_de_edificios(client):
    escalas = client.get(f"/api/escalas?convenio={SUTERYH}").json
    assert len(escalas) == 3 * len(nombres_escala())
    assert not any(e["verificada"] for e in escalas)
    pagina = client.get(f"/escalas?convenio={SUTERYH}").text
    assert "2ª cat." in pagina and "Valor vivienda" in pagina and "1.214.974,00" in pagina


def test_sueldo_de_encargado_desde_la_web(client):
    e = consorcio(client)
    r = encargado(client, e)
    assert "guardado" in r.text and 'id="edificio"' in r.text
    emp = client.get(f"/api/empresas/{e}/empleados").json[0]
    assert emp["extras"] == {"afiliado": True, "retira_residuos": True, "tareas": [], "tramos_titulo": 0}

    # Sin los datos del edificio no se puede liquidar.
    r = client.post(f"/empresas/{e}/liquidar", data={**PAGO, "tipo": "mensual", "periodo": "2026-09"},
                    follow_redirects=True)
    assert "Faltan los datos del edificio" in r.text

    client.post(f"/empresas/{e}/edificio", data={"categoria": "2", "unidades_funcionales": "20"})
    form = client.get(f"/empresas/{e}/liquidar?periodo=2026-09&tipo=mensual").text
    assert f'name="horas_50_{emp["id"]}"' in form and f'name="inasistencias_{emp["id"]}"' not in form
    r = client.post(f"/empresas/{e}/liquidar", data={**PAGO, "tipo": "mensual", "periodo": "2026-09"},
                    follow_redirects=True)
    assert "Listo: 1 recibos" in r.text and "1.243.751,00" in r.text
    liq_id = re.search(r'href="/liquidaciones/(\d+)"', r.text).group(1)
    detalle = client.get(f"/liquidaciones/{liq_id}").text
    assert "Contribuciones del empleador del convenio" in detalle and "63.578,28" in detalle

    # Con horas extra por empleado, y la pantalla del recibo muestra las contribuciones del convenio.
    r = client.post(f"/empresas/{e}/liquidar", data={**PAGO, "tipo": "mensual", "periodo": "2026-09",
                                                     f"horas_50_{emp['id']}": "10"}, follow_redirects=True)
    assert "Listo: 1 recibos" in r.text and "1.243.751,00" not in r.text
    pdf = PdfReader(BytesIO(client.get(f"/api/empresas/{e}/recibos/2026-09.pdf").data))
    texto = "".join(p.extract_text() for p in pdf.pages)
    assert "Horas extra 50%" in texto and "Vivienda (en especie)" not in texto and "CCT 589/10" in texto
    assert client.get(f"/api/empresas/{e}/empleados").json[0]["jornada_horas"] == 8


def test_datos_del_encargado_y_errores(client):
    e = consorcio(client, categoria="1", unidades_funcionales="10", zona_desfavorable="0")
    data = {"apellido": "Gómez", "nombre": "Juan", "cuil": "20-22222222-3", "convenio": SUTERYH,
            "categoria": "Ayudante Media jornada", "fecha_ingreso": "2020-01-01", "afiliado": "0",
            "tramos_titulo": "1"}
    r = client.post(f"/empresas/{e}/empleados", data={**data, "tareas": ["jardin", "viaticos"]},
                    follow_redirects=True)
    assert "guardado" in r.text
    emp = client.get(f"/api/empresas/{e}/empleados").json[0]
    assert emp["jornada_horas"] == 4
    assert emp["extras"] == {"afiliado": False, "retira_residuos": False, "tareas": ["jardin", "viaticos"],
                             "tramos_titulo": 1}
    # El formulario de edición los muestra cargados.
    pagina = client.get(f"/empresas/{e}?editar={emp['id']}").text
    assert 'value="jardin" checked' in pagina and 'value="limpieza_piletas">' in pagina

    r = client.post("/api/liquidaciones", json={**PAGO, "empleado_id": emp["id"], "periodo": "2026-09",
                                                "horas_100": "2"})
    assert r.status_code == 201
    codigos = {c["codigo"] for c in r.json["conceptos"]}
    assert {"JARD", "VIAT", "TIT", "HE100"} <= codigos and "SIND" not in codigos and "RES" not in codigos

    r = client.post("/api/liquidaciones", json={**PAGO, "empleado_id": emp["id"], "periodo": "2026-09",
                                                "inasistencias_injustificadas": "1"})
    assert r.status_code == 400 and "faltas" in r.json["error"]
    r = client.post("/api/liquidaciones/final", json={**PAGO, "empleado_id": emp["id"], "causa": "renuncia",
                                                      "fecha_egreso": "2026-10-15"})
    assert r.status_code == 400 and "todavía no está programada" in r.json["error"]

    r = client.post("/api/empleados", json={**data, "empresa_id": e, "cuil": "20-11111111-2",
                                            "tareas": ["lavar autos"]})
    assert r.status_code == 400 and "Tarea desconocida" in r.json["error"]
    r = client.put(f"/api/empresas/{e}/edificio", json={"categoria": 7})
    assert r.status_code == 400
    r = client.put(f"/api/empresas/{e}/edificio", json={"categoria": 3, "unidades_funcionales": 12,
                                                         "zona_desfavorable": True})
    assert r.json == {"empresa_id": e, "categoria": 3, "unidades_funcionales": 12, "zona_desfavorable": True,
                      "zona_base": "remunerativo", "zona_recibo_aparte": False}
    r = client.put(f"/api/empresas/{e}/edificio", json={"categoria": 3, "zona_base": "todo"})
    assert r.status_code == 400 and "zona_base" in r.json["error"]


def test_zona_fria_en_recibo_aparte(client):
    """Mismo caso que el recibo real de vigilancia nocturna, liquidado desde la web."""
    e = consorcio(client, categoria="3", zona_desfavorable="1", zona_base="remunerativo", zona_recibo_aparte="1")
    encargado(client, e, categoria="Personal Vigilancia Nocturna", fecha_ingreso="2015-09-01", afiliado="0",
              retira_residuos="0", fecha_egreso="2026-09-30")
    r = client.post(f"/empresas/{e}/liquidar", data={**PAGO, "tipo": "mensual", "periodo": "2026-09"},
                    follow_redirects=True)
    assert "Listo: 1 recibos de 2026-09, más 1 de zona fría aparte." in r.text and "1.320.544,00" in r.text
    zona = client.get(f"/empresas/{e}/liquidaciones/2026-09?tipo=zona_fria").text
    assert "Zona fría 2026-09" in zona and "660.272,00" in zona
    pdf = PdfReader(BytesIO(client.get(f"/api/empresas/{e}/recibos/2026-09.pdf?tipo=zona_fria").data))
    assert "zona fría" in pdf.pages[0].extract_text()
    assert "Zona fría" in client.get(f"/empresas/{e}").text
    assert client.get("/").status_code == 200

    # La zona del recibo aparte tiene su propio aguinaldo, en otro recibo (como la hoja SAC de la planilla).
    r = client.post(f"/empresas/{e}/liquidar", data={**PAGO, "tipo": "sac", "periodo": "2026-09"},
                    follow_redirects=True)
    assert "más 1 de zona fría aparte" in r.text
    sac_id = re.search(r'href="/liquidaciones/(\d+)"', r.text).group(1)
    sac = client.get(f"/liquidaciones/{sac_id}").text
    assert "1.645.537,30" in sac and "822.768,65" not in sac
    sac_zona = client.get(f"/empresas/{e}/liquidaciones/2026-09?tipo=sac_zona_fria").text
    assert "Aguinaldo s/ zona fría 2026-09" in sac_zona
    liq_id = re.search(r'href="/liquidaciones/(\d+)"', sac_zona).group(1)
    detalle = client.get(f"/liquidaciones/{liq_id}").text
    assert "SAC s/ zona fría" in detalle and "822.768,65" in detalle
    pdf = PdfReader(BytesIO(client.get(f"/api/empresas/{e}/recibos/2026-09.pdf?tipo=sac_zona_fria").data))
    assert "SAC s/ zona fría" in pdf.pages[0].extract_text()

    # Si el consorcio pasa la zona al mismo recibo, el recibo aparte se borra al reliquidar.
    client.post(f"/empresas/{e}/edificio", data={"categoria": "3", "zona_desfavorable": "1"})
    client.post(f"/empresas/{e}/liquidar", data={**PAGO, "tipo": "mensual", "periodo": "2026-09"})
    assert client.get(f"/api/empresas/{e}/recibos/2026-09.pdf?tipo=zona_fria").status_code == 404


def test_aguinaldo_de_encargado(client):
    e = consorcio(client, categoria="2", unidades_funcionales="20")
    encargado(client, e, fecha_egreso="2026-09-30")
    for periodo in ("2026-07", "2026-08", "2026-09"):
        client.post(f"/empresas/{e}/liquidar", data={**PAGO, "tipo": "mensual", "periodo": periodo})
    r = client.post(f"/api/empresas/{e}/sac", json={**PAGO, "periodo": "2026-09"})
    assert r.status_code == 201 and r.json["liquidaciones"][0]["neto"] == "310938.00"


def test_plantilla_de_escala_de_edificios(client):
    plantilla = load_workbook(BytesIO(client.get(f"/api/escalas/plantilla?convenio={SUTERYH}").data))
    filas = [r[0].value for r in plantilla["Escala"].iter_rows(min_row=2)]
    assert filas == nombres_escala() and "Intendente|3" in filas and "Valor vivienda" in filas
    nueva = generar_plantilla({n: D("1000") for n in nombres_escala()}, vigencia=date(2026, 10, 1),
                              categorias=nombres_escala())
    r = client.post("/escalas/importar", data={"convenio": SUTERYH, "archivo": (BytesIO(nueva), "oct.xlsx")},
                    follow_redirects=True)
    assert f"Se cargaron {len(nombres_escala())} filas" in r.text and "2026-10-01" in r.text


def test_base_vieja_suma_columna_extras_y_motor(tmp_path):
    ruta = tmp_path / "vieja.db"
    conn = sqlite3.connect(ruta)
    conn.executescript("""
        CREATE TABLE convenios (codigo TEXT PRIMARY KEY, nombre TEXT NOT NULL, tiene_motor INTEGER NOT NULL DEFAULT 0);
        INSERT INTO convenios VALUES ('CCT 589/10', 'Cargado a mano', 0);
        CREATE TABLE empleados (id INTEGER PRIMARY KEY AUTOINCREMENT, empresa_id INTEGER NOT NULL,
            legajo TEXT, apellido TEXT NOT NULL, nombre TEXT NOT NULL, cuil TEXT NOT NULL,
            convenio TEXT NOT NULL DEFAULT 'CCT 130/75', categoria TEXT NOT NULL, fecha_ingreso TEXT NOT NULL,
            fecha_egreso TEXT, jornada_horas INTEGER NOT NULL DEFAULT 8);
    """)
    conn.close()
    client = create_app({"DATABASE": str(ruta), "TESTING": True}).test_client()
    conv = {c["codigo"]: c for c in client.get("/api/convenios").json}
    assert conv[SUTERYH]["tiene_motor"] is True
    conn = dbmod.connect(str(ruta))
    assert "extras" in {r["name"] for r in conn.execute("PRAGMA table_info(empleados)")}
    dbmod.init_schema(conn)  # idempotente
