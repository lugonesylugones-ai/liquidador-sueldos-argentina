"""Alta masiva de empleados desde una plantilla Excel.

Hoja "Empleados":
    Legajo | Apellido | Nombre | CUIL | Convenio | Categoría | Fecha ingreso | Jornada (hs) | Fecha egreso

Legajo, Convenio (default CCT 130/75), Jornada (default 8) y Fecha egreso son opcionales.
Si una fila tiene errores no se carga ninguna.
"""
from dataclasses import dataclass
from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .cuit import normalizar_cuit
from .escalas import CONVENIO_COMERCIO, _parse_fecha, normalizar_categoria

HOJA = "Empleados"
COLUMNAS = ("Legajo", "Apellido", "Nombre", "CUIL", "Convenio", "Categoría", "Fecha ingreso",
            "Jornada (hs)", "Fecha egreso")


@dataclass
class FilaEmpleado:
    legajo: str | None
    apellido: str
    nombre: str
    cuil: str
    convenio: str
    categoria: str
    fecha_ingreso: date
    jornada_horas: int
    fecha_egreso: date | None


def generar_plantilla_empleados(categorias: dict) -> bytes:
    """`categorias` = {convenio: [categoría, ...]} para la hoja de ayuda."""
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    ws.append(COLUMNAS)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
    for col, ancho in zip("ABCDEFGHI", (10, 20, 20, 16, 14, 26, 14, 12, 14)):
        ws.column_dimensions[col].width = ancho
    ayuda = wb.create_sheet("Categorías")
    ayuda.append(["Convenio", "Categoría"])
    for convenio, cats in categorias.items():
        for cat in cats:
            ayuda.append([convenio, cat])
    ayuda.column_dimensions["A"].width = 14
    ayuda.column_dimensions["B"].width = 28
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def validar_categoria(convenio: str, texto: str, categorias: dict) -> str:
    """Devuelve el nombre canónico de la categoría o lanza ValueError."""
    if convenio not in categorias:
        raise ValueError(f"convenio desconocido {convenio!r}")
    validas = categorias[convenio]
    if convenio == CONVENIO_COMERCIO:
        cat = normalizar_categoria(texto)
    else:
        cat = next((c for c in validas if c.lower() == texto.strip().lower()), None)
    if cat is None or cat not in validas:
        raise ValueError(f"categoría {texto!r} no existe en {convenio}")
    return cat


def leer_empleados(contenido: bytes, categorias: dict) -> tuple[list, list]:
    """Devuelve (filas, errores)."""
    filas, errores = [], []
    try:
        wb = load_workbook(BytesIO(contenido), data_only=True)
    except Exception as exc:
        return [], [f"No se pudo abrir el Excel: {exc}"]
    if HOJA not in wb.sheetnames:
        return [], [f"Falta la hoja '{HOJA}'"]
    ws = wb[HOJA]
    encabezado = tuple(_texto(c.value) for c in ws[1][:len(COLUMNAS)])
    if encabezado != COLUMNAS:
        return [], [f"Encabezado esperado {COLUMNAS}, se encontró {encabezado}"]
    cuiles, legajos = set(), set()
    for n, fila in enumerate(ws.iter_rows(min_row=2, max_col=len(COLUMNAS), values_only=True), start=2):
        fila = tuple(fila) + (None,) * (len(COLUMNAS) - len(fila))
        if all(_texto(v) == "" for v in fila):
            continue
        legajo, apellido, nombre, cuil, convenio, categoria, ingreso, jornada, egreso = fila
        err = []
        apellido, nombre = _texto(apellido), _texto(nombre)
        if not apellido or not nombre:
            err.append("falta apellido o nombre")
        try:
            cuil = normalizar_cuit(cuil)
            if cuil in cuiles:
                err.append(f"CUIL {cuil} repetido")
            cuiles.add(cuil)
        except ValueError as exc:
            err.append(str(exc))
        convenio = _texto(convenio) or CONVENIO_COMERCIO
        try:
            categoria = validar_categoria(convenio, _texto(categoria), categorias)
        except ValueError as exc:
            err.append(str(exc))
        try:
            ingreso = _parse_fecha(ingreso)
        except ValueError as exc:
            err.append(f"ingreso: {exc}")
        try:
            jornada = int(float(_texto(jornada) or 8))
            if not 1 <= jornada <= 8:
                raise ValueError
        except ValueError:
            err.append(f"jornada inválida {jornada!r} (1 a 8 hs)")
        if _texto(egreso):
            try:
                egreso = _parse_fecha(egreso)
            except ValueError as exc:
                err.append(f"egreso: {exc}")
        else:
            egreso = None
        legajo = _texto(legajo) or None
        if legajo:
            if legajo in legajos:
                err.append(f"legajo {legajo} repetido")
            legajos.add(legajo)
        if err:
            errores.append(f"Fila {n}: " + "; ".join(err))
        else:
            filas.append(FilaEmpleado(legajo, apellido, nombre, cuil, convenio, categoria,
                                      ingreso, jornada, egreso))
    if not filas and not errores:
        errores.append("La plantilla no tiene empleados")
    return filas, errores
