"""Genera una escala de ejemplo y un recibo de punta a punta usando la API.

    python scripts/generar_ejemplo.py

Deja en ejemplos/ las escalas de julio a septiembre 2026 cargadas en la
plantilla y el recibo PDF (septiembre).
"""
import sys
import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from backend.app import create_app  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from backend.escalas import HOJA, agregar_fila, generar_plantilla  # noqa: E402

SALIDA = RAIZ / "ejemplos"

# Escalas julio a septiembre 2026 (acuerdo julio 2026) según Ignacio Online, no la
# circular de FAECYS. No remunerativo = Inc. NR 100.000 + Recomp. NR 20.000; julio y
# agosto traen además una asignación de única vez de 25.000. Auxiliar B coincide con
# los recibos reales de esos tres meses; el resto de las categorías no está contrastado.
CATEGORIAS_ORDEN = [
    "Maestranza A", "Maestranza B", "Maestranza C",
    "Administrativo A", "Administrativo B", "Administrativo C",
    "Administrativo D", "Administrativo E", "Administrativo F",
    "Cajero A", "Cajero B", "Cajero C", "Auxiliar A", "Auxiliar B", "Auxiliar C",
    "Auxiliar Especializado A", "Auxiliar Especializado B",
    "Vendedor A", "Vendedor B", "Vendedor C", "Vendedor D",
]
BASICOS_2026 = {
    date(2026, 7, 1): [1137023, 1140294, 1151751, 1149298, 1154212, 1159120, 1173854, 1186128,
                       1204135, 1153389, 1159120, 1166487, 1153389, 1161573, 1188584, 1163214,
                       1177944, 1153389, 1177947, 1186128, 1204135],
    date(2026, 8, 1): [1160461, 1163793, 1175464, 1172965, 1177971, 1182970, 1197978, 1210482,
                       1228824, 1177132, 1182970, 1190474, 1177132, 1185469, 1212983, 1187140,
                       1202145, 1177132, 1202148, 1210482, 1228824],
    date(2026, 9, 1): [1183900, 1187292, 1199176, 1196632, 1201729, 1206821, 1222103, 1234836,
                       1253514, 1200875, 1206821, 1214462, 1200875, 1209365, 1237383, 1211067,
                       1226346, 1200875, 1226349, 1234836, 1253514],
}
NO_REM_2026 = 120000
ASIG_UNICA_2026 = {date(2026, 7, 1): 25000, date(2026, 8, 1): 25000}


def escala_2026() -> bytes:
    wb = load_workbook(BytesIO(generar_plantilla()))
    ws = wb[HOJA]
    ws.delete_rows(2, ws.max_row)
    for vigencia, basicos in BASICOS_2026.items():
        for cat, basico in zip(CATEGORIAS_ORDEN, basicos):
            agregar_fila(ws, cat, vigencia, basico, NO_REM_2026, ASIG_UNICA_2026.get(vigencia))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> Path:
    SALIDA.mkdir(exist_ok=True)
    plantilla = escala_2026()
    (SALIDA / "escala_comercio_2026_jul_sep.xlsx").write_bytes(plantilla)

    with tempfile.TemporaryDirectory() as tmp:
        app = create_app({"DATABASE": str(Path(tmp) / "ejemplo.db"), "TESTING": True})
        c = app.test_client()
        r = c.post("/escalas/importar", data={"archivo": (BytesIO(plantilla), "escala.xlsx")},
                   content_type="multipart/form-data")
        assert r.status_code == 201, r.json
        empresa = c.post("/empresas", json={
            "razon_social": "Almacén Ejemplo S.R.L.", "cuit": "30-71234567-8",
            "domicilio": "Av. Corrientes 1234, CABA"}).json["id"]
        empleado = c.post("/empleados", json={
            "empresa_id": empresa, "apellido": "Pérez", "nombre": "María Laura",
            "cuil": "27-30123456-4", "categoria": "Auxiliar B",
            "fecha_ingreso": "2017-07-03", "jornada_horas": 8}).json["id"]
        r = c.post("/liquidaciones", json={
            "empleado_id": empleado, "periodo": "2026-09", "inasistencias_injustificadas": 0,
            "fecha_pago": "2026-10-03", "lugar_pago": "Ciudad Autónoma de Buenos Aires",
            "ultimo_deposito_periodo": "08/2026", "ultimo_deposito_fecha": "2026-09-10",
            "ultimo_deposito_banco": "Banco de la Nación Argentina"})
        assert r.status_code == 201, r.json
        pdf = c.get(f"/liquidaciones/{r.json['id']}/recibo.pdf")
        assert pdf.status_code == 200
        destino = SALIDA / "recibo_ejemplo.pdf"
        destino.write_bytes(pdf.data)
    print(f"Recibo generado en {destino}")
    return destino


if __name__ == "__main__":
    main()
